import os
import time
import shutil
import subprocess
import io
import csv
import clickhouse_connect
import logging
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import ipaddress
import config

logger = logging.getLogger("collector")

class MetricsCollector:
    def __init__(self):
        self.client = self._get_client()
        self._cache = {}
        self._cache_time = {}

    def _get_client(self):
        return clickhouse_connect.get_client(
            host=config.CLICKHOUSE_HOST,
            port=config.CLICKHOUSE_PORT,
            username=config.CLICKHOUSE_USER,
            password=config.CLICKHOUSE_PASSWORD,
            database=config.CLICKHOUSE_DB
        )

    def _query(self, sql: str, parameters=None, settings=None):
        if settings is None:
            settings = {
                "timeout_before_checking_execution_speed": 0,
                "max_execution_time": 45,
                "max_threads": 8
            }
        client = self._get_client()
        try:
            return client.query(sql, parameters=parameters, settings=settings)
        finally:
            client.close()

    def get_system_health(self):
        try:
            with open('/proc/uptime', 'r') as f:
                uptime_sec = float(f.read().split()[0])
            days = int(uptime_sec // 86400)
            hours = int((uptime_sec % 86400) // 3600)
            minutes = int((uptime_sec % 3600) // 60)
            uptime_str = f"{days}d {hours}h {minutes}m"
        except Exception:
            uptime_str = "Unknown"
            uptime_sec = 0

        try:
            load1, load5, load15 = os.getloadavg()
        except Exception:
            load1 = load5 = load15 = 0.0
            
        cpu_cores = os.cpu_count() or 1
        load_pct = round((load1 / cpu_cores) * 100, 1)

        meminfo = {}
        try:
            with open('/proc/meminfo', 'r') as f:
                for line in f:
                    parts = line.split(':')
                    if len(parts) == 2:
                        key = parts[0].strip()
                        val = parts[1].strip().split()[0]
                        meminfo[key] = int(val) * 1024 # in bytes
        except Exception:
            pass

        mem_total = meminfo.get('MemTotal', 0)
        mem_avail = meminfo.get('MemAvailable', 0)
        mem_used = mem_total - mem_avail
        mem_pct = round((mem_used / mem_total) * 100, 1) if mem_total else 0

        try:
            disk_root = shutil.disk_usage('/')
            disk_ch = shutil.disk_usage('/var/lib/clickhouse') if os.path.exists('/var/lib/clickhouse') else disk_root
        except Exception:
            disk_root = disk_ch = None

        services = {}
        for svc in ['clickhouse-server', 'vector', 'nat-ai-agent']:
            try:
                res = subprocess.run(['systemctl', 'is-active', svc], capture_output=True, text=True, timeout=2)
                services[svc] = (res.stdout.strip() == 'active')
            except Exception:
                services[svc] = False

        return {
            "uptime": uptime_str,
            "uptime_seconds": uptime_sec,
            "cpu_cores": cpu_cores,
            "load_avg": {
                "1m": round(load1, 2),
                "5m": round(load5, 2),
                "15m": round(load15, 2)
            },
            "load_pct": min(100.0, load_pct),
            "memory": {
                "total_bytes": mem_total,
                "used_bytes": mem_used,
                "free_bytes": mem_avail,
                "pct": mem_pct
            },
            "disk": {
                "root_total": disk_root.total if disk_root else 0,
                "root_used": disk_root.used if disk_root else 0,
                "root_free": disk_root.free if disk_root else 0,
                "root_pct": round((disk_root.used / disk_root.total) * 100, 1) if disk_root else 0,
                "ch_free": disk_ch.free if disk_ch else 0
            },
            "services": services
        }

    def get_storage_metrics(self):
        try:
            q = """
            SELECT 
                sum(rows) AS total_rows,
                sum(data_compressed_bytes) AS compressed_bytes,
                sum(data_uncompressed_bytes) AS uncompressed_bytes
            FROM system.parts
            WHERE active AND table='translations'
            """
            r = self._query(q).result_rows[0]
            total_rows = r[0] or 0
            comp_bytes = r[1] or 0
            uncomp_bytes = r[2] or 0
            ratio = round(uncomp_bytes / comp_bytes, 2) if comp_bytes else 0
            bytes_per_row = round(comp_bytes / total_rows, 2) if total_rows else 0
        except Exception as e:
            logger.error(f"Error querying storage metrics: {e}")
            total_rows = comp_bytes = uncomp_bytes = ratio = bytes_per_row = 0

        return {
            "total_rows": total_rows,
            "compressed_bytes": comp_bytes,
            "uncompressed_bytes": uncomp_bytes,
            "compression_ratio": ratio,
            "bytes_per_row": bytes_per_row,
            "retention_days": 180,
            "sorting_key": "(toStartOfHour(timestamp), router_ip, nat_src_ip, nat_src_port, timestamp)",
            "codec": "DoubleDelta, T64, ZSTD(12)"
        }

    def _build_telemetry_time_filter(self, minutes=None, start_time=None, end_time=None, tz="Asia/Kathmandu"):
        params = {"tz": tz}
        
        if start_time and end_time:
            st = str(start_time).replace('T', ' ').strip()
            et = str(end_time).replace('T', ' ').strip()
            if len(st) == 10: st += ' 00:00:00'
            elif len(st) == 16: st += ':00'
            if len(et) == 10: et += ' 23:59:59'
            elif len(et) == 16: et += ':59'
            
            params['st'] = st
            params['et'] = et
            
            try:
                dt_start = datetime.strptime(st[:19], "%Y-%m-%d %H:%M:%S")
                dt_end = datetime.strptime(et[:19], "%Y-%m-%d %H:%M:%S")
                window_sec = max(1.0, (dt_end - dt_start).total_seconds())
            except Exception:
                window_sec = 3600.0

            where_sql = f"""
            WHERE toStartOfHour(timestamp) >= toStartOfHour(toDateTime64(%(st)s, 3, '{tz}'))
              AND toStartOfHour(timestamp) <= toStartOfHour(toDateTime64(%(et)s, 3, '{tz}'))
              AND timestamp >= toDateTime64(%(st)s, 3, '{tz}')
              AND timestamp <= toDateTime64(%(et)s, 3, '{tz}')
            """
            mins = int(window_sec // 60)
            if mins < 60:
                label = f"Custom ({mins}m)"
            elif mins < 1440:
                label = f"Custom ({round(mins/60, 1)}h)"
            else:
                label = f"Custom ({round(mins/1440, 1)}d)"

            return {
                "where_sql": where_sql,
                "params": params,
                "window_sec": window_sec,
                "window_mins": mins,
                "window_label": label,
                "is_custom": True,
                "start_time": st,
                "end_time": et
            }
        
        # Relative range
        mins = int(minutes or 5)
        window_sec = max(1.0, mins * 60.0)
        
        if mins < 60:
            label = f"Last {mins} Minutes"
        elif mins == 60:
            label = "Last 1 Hour"
        elif mins < 1440:
            label = f"Last {mins // 60} Hours"
        elif mins == 1440:
            label = "Last 24 Hours"
        elif mins == 2880:
            label = "Last 2 Days (48h)"
        elif mins == 4320:
            label = "Last 3 Days (72h)"
        else:
            label = f"Last {round(mins/1440, 1)} Days"

        where_sql = f"""
        WHERE toStartOfHour(timestamp) >= toStartOfHour(now64(3, '{tz}') - INTERVAL {mins} MINUTE)
          AND timestamp >= now64(3, '{tz}') - INTERVAL {mins} MINUTE
        """
        
        return {
            "where_sql": where_sql,
            "params": params,
            "window_sec": window_sec,
            "window_mins": mins,
            "window_label": label,
            "is_custom": False,
            "start_time": None,
            "end_time": None
        }

    def get_ingestion_telemetry(self, time_ctx):
        where_sql = time_ctx["where_sql"]
        params = time_ctx["params"]
        window_sec = time_ctx["window_sec"]
        window_mins = time_ctx["window_mins"]
        tz = params.get("tz", "Asia/Kathmandu")
        
        try:
            # 1. EPS last 15s for instantaneous pulse
            q_rate = "SELECT count() as cnt, round(count() / 15.0, 1) as eps FROM nat_logs.translations WHERE toStartOfHour(timestamp) >= toStartOfHour(now() - INTERVAL 1 MINUTE) AND timestamp >= now() - INTERVAL 15 SECOND"
            r_rate = self._query(q_rate).result_rows[0]
            current_eps = r_rate[1]

            # For large windows (> 60m), use pre-aggregated minute_rollups for instant (<1s) responses
            if window_mins > 60:
                if time_ctx["is_custom"]:
                    rollup_where = f"WHERE minute >= toDateTime(%(st)s, '{tz}') AND minute <= toDateTime(%(et)s, '{tz}')"
                else:
                    rollup_where = f"WHERE minute >= now('{tz}') - INTERVAL {window_mins} MINUTE"

                q_routers = f"""
                SELECT 
                    router_ip,
                    sum(total_count) AS flows,
                    round(sum(total_count) / {window_sec}, 1) AS eps,
                    max(minute) AS last_seen,
                    uniqMerge(src_ip_state) AS subscribers,
                    uniqMerge(nat_src_ip_state) AS nat_ips
                FROM nat_logs.minute_rollups
                {rollup_where}
                GROUP BY router_ip
                ORDER BY flows DESC
                """
                r_routers = self._query(q_routers, parameters=params).result_rows
                routers_list = []
                for r in r_routers:
                    routers_list.append({
                        "router_ip": str(r[0]),
                        "flows": int(r[1]),
                        "flows_5m": int(r[1]),
                        "eps": float(r[2]),
                        "last_seen": str(r[3]),
                        "subscribers": int(r[4]),
                        "subscribers_5m": int(r[4]),
                        "nat_ips": int(r[5]),
                        "nat_ips_5m": int(r[5])
                    })

                q_proto = f"""
                SELECT 
                    protocol,
                    sum(total_count) AS flows,
                    round(sum(total_count) * 100.0 / sum(sum(total_count)) OVER(), 1) AS pct
                FROM nat_logs.minute_rollups
                {rollup_where}
                GROUP BY protocol
                ORDER BY flows DESC
                """
                r_proto = self._query(q_proto, parameters=params).result_rows
                protocols = [{"protocol": str(r[0]), "flows": int(r[1]), "pct": float(r[2])} for r in r_proto]

                q_ports = f"""
                SELECT 
                    dst_port,
                    topK(1)(protocol)[1] as proto,
                    sum(total_count) AS flows
                FROM nat_logs.minute_rollups
                {rollup_where}
                GROUP BY dst_port
                ORDER BY flows DESC
                LIMIT 8
                """
                r_ports = self._query(q_ports, parameters=params).result_rows
                top_ports = [{"port": int(r[0]), "proto": str(r[1]), "flows": int(r[2])} for r in r_ports]

            else:
                # 2. Routers in live window (<= 60m)
                q_routers = f"""
                SELECT 
                    router_ip,
                    count() AS flows,
                    round(count() / {window_sec}, 1) AS eps,
                    max(timestamp) AS last_seen,
                    uniq(src_ip) AS subscribers,
                    uniq(nat_src_ip) AS nat_ips
                FROM nat_logs.translations
                {where_sql}
                GROUP BY router_ip
                ORDER BY flows DESC
                """
                r_routers = self._query(q_routers, parameters=params).result_rows
                routers_list = []
                for r in r_routers:
                    routers_list.append({
                        "router_ip": str(r[0]),
                        "flows": r[1],
                        "flows_5m": r[1],
                        "eps": r[2],
                        "last_seen": str(r[3]),
                        "subscribers": r[4],
                        "subscribers_5m": r[4],
                        "nat_ips": r[5],
                        "nat_ips_5m": r[5]
                    })

                # 3. Protocols in window
                q_proto = f"""
                SELECT 
                    protocol,
                    count() AS flows,
                    round(count() * 100.0 / sum(count()) OVER(), 1) AS pct
                FROM nat_logs.translations
                {where_sql}
                GROUP BY protocol
                ORDER BY flows DESC
                """
                r_proto = self._query(q_proto, parameters=params).result_rows
                protocols = [{"protocol": str(r[0]), "flows": r[1], "pct": r[2]} for r in r_proto]

                # 4. Top Ports in window
                q_ports = f"""
                SELECT 
                    dst_port,
                    topK(1)(protocol)[1] as proto,
                    count() AS flows
                FROM nat_logs.translations
                {where_sql}
                GROUP BY dst_port
                ORDER BY flows DESC
                LIMIT 8
                """
                r_ports = self._query(q_ports, parameters=params).result_rows
                top_ports = [{"port": r[0], "proto": str(r[1]), "flows": r[2]} for r in r_ports]

            return {
                "current_eps": current_eps,
                "window_eps": round(sum(r["flows"] for r in routers_list) / window_sec, 1) if routers_list else 0,
                "active_routers_count": len(routers_list),
                "routers": routers_list,
                "protocols": protocols,
                "top_ports": top_ports,
                "transports": [
                    {"protocol": "Syslog (RFC5424 / RFC3164)", "port": "UDP 514", "status": "Active Streaming"},
                    {"protocol": "IPFIX / NetFlow v9", "port": "UDP 2055", "status": "Active Streaming"}
                ]
            }
        except Exception as e:
            logger.error(f"Error querying ingestion telemetry: {e}")
            return {
                "current_eps": 0,
                "window_eps": 0,
                "active_routers_count": 0,
                "routers": [],
                "protocols": [],
                "top_ports": [],
                "transports": []
            }

    def get_public_nat_pool_details(self, time_ctx, limit=500, total_nat_ips_hint=None):
        where_sql = time_ctx["where_sql"]
        params = time_ctx["params"]
        window_mins = time_ctx["window_mins"]
        tz = params.get("tz", "Asia/Kathmandu")
        
        try:
            total_public_ips = total_nat_ips_hint
            
            # For large windows, inspect active 1-hour partition to get fast live port utilization
            if window_mins > 60:
                scan_where = f"WHERE toStartOfHour(timestamp) >= toStartOfHour(now64(3, '{tz}') - INTERVAL 1 HOUR) AND timestamp >= now64(3, '{tz}') - INTERVAL 1 HOUR"
            else:
                scan_where = where_sql

            if total_public_ips is None:
                q_nat_count = f"SELECT uniq(nat_src_ip) FROM nat_logs.translations {scan_where}"
                r_nat_count = self._query(q_nat_count, parameters=params).result_rows
                total_public_ips = r_nat_count[0][0] if r_nat_count else 0

            q_nat_pool = f"""
            SELECT 
                nat_src_ip,
                router_ip,
                uniqExact(nat_src_port) AS active_ports,
                round((uniqExact(nat_src_port) / 64512.0) * 100, 1) AS util_pct,
                count() AS total_flows,
                uniqExact(src_ip) AS active_subscribers
            FROM nat_logs.translations
            {scan_where}
            GROUP BY nat_src_ip, router_ip
            ORDER BY active_ports DESC, util_pct DESC, total_flows DESC
            LIMIT {int(limit)}
            """
            r_nat = self._query(q_nat_pool, parameters=params).result_rows
            nat_pool_list = []
            for r in r_nat:
                nat_pool_list.append({
                    "nat_src_ip": str(r[0]),
                    "router_ip": str(r[1]),
                    "active_ports": r[2],
                    "util_pct": r[3],
                    "total_flows": r[4],
                    "active_subscribers": r[5],
                    "status": "CRITICAL" if r[3] >= 85 else ("WARNING" if r[3] >= 70 else "HEALTHY")
                })

            nat_pool_list.sort(key=lambda x: (x.get("active_ports", 0), x.get("util_pct", 0), x.get("total_flows", 0)), reverse=True)
            return {
                "total_public_nat_ips": total_public_ips,
                "top_nat_ips": nat_pool_list
            }
        except Exception as e:
            logger.error(f"Error querying public NAT pool: {e}")
            return {"total_public_nat_ips": 0, "top_nat_ips": []}

    def get_private_pool_details(self, time_ctx, limit=100, total_flows_hint=None, total_subs_hint=None):
        where_sql = time_ctx["where_sql"]
        params = time_ctx["params"]
        window_mins = time_ctx["window_mins"]
        tz = params.get("tz", "Asia/Kathmandu")
        
        try:
            total_private_ips = total_subs_hint
            total_flows = total_flows_hint

            # For large windows (> 60m), analyze top active subscribers from recent 1-hour window for speed
            if window_mins > 60:
                scan_where = f"WHERE toStartOfHour(timestamp) >= toStartOfHour(now64(3, '{tz}') - INTERVAL 1 HOUR) AND timestamp >= now64(3, '{tz}') - INTERVAL 1 HOUR"
            else:
                scan_where = where_sql

            if total_flows is None or total_private_ips is None:
                q_priv_count = f"SELECT uniq(src_ip) AS total_private_subscribers, count() AS total_flows FROM nat_logs.translations {scan_where}"
                r_priv = self._query(q_priv_count, parameters=params).result_rows[0]
                total_private_ips = r_priv[0] or 0
                total_flows = r_priv[1] or 0

            q_top_sub = f"""
            SELECT 
                src_ip,
                router_ip,
                count() AS flows,
                uniq(dst_ip) AS distinct_dests,
                uniq(dst_port) AS distinct_ports,
                topK(1)(nat_src_ip)[1] AS assigned_nat_ip
            FROM nat_logs.translations
            {scan_where}
            GROUP BY src_ip, router_ip
            ORDER BY distinct_ports DESC, flows DESC
            LIMIT {int(limit)}
            """
            r_top_sub = self._query(q_top_sub, parameters=params).result_rows
            top_subscribers = []
            
            heavy_thresh = max(20000, min(window_mins, 60) * 4000)
            scanner_thresh = max(300, min(2000, min(window_mins, 60) * 15))

            for r in r_top_sub:
                flows = r[2]
                distinct_dests = r[3]
                distinct_ports = r[4]
                is_scanner = (distinct_dests >= scanner_thresh) or (distinct_ports >= 1000)
                if distinct_dests >= scanner_thresh and distinct_ports >= 1000:
                    behavior = "SCANNER / BOTNET"
                elif distinct_ports >= 1000:
                    behavior = "PORT SWEEPER (SCANNER)"
                elif distinct_dests >= scanner_thresh:
                    behavior = "SCANNER / FLOOD"
                elif flows > heavy_thresh:
                    behavior = "HEAVY"
                else:
                    behavior = "NORMAL"

                top_subscribers.append({
                    "src_ip": str(r[0]),
                    "router_ip": str(r[1]),
                    "flows": flows,
                    "distinct_dests": distinct_dests,
                    "distinct_ports": distinct_ports,
                    "assigned_nat_ip": str(r[5]),
                    "behavior": behavior
                })

            return {
                "total_private_subscribers": total_private_ips,
                "total_flows": total_flows,
                "total_flows_5m": total_flows,
                "top_subscribers": top_subscribers
            }
        except Exception as e:
            logger.error(f"Error querying private pool: {e}")
            return {"total_private_subscribers": 0, "total_flows": 0, "total_flows_5m": 0, "top_subscribers": []}

    def get_full_telemetry(self, minutes=5, start_time=None, end_time=None):
        tz = getattr(config, 'SERVER_TIMEZONE', 'Asia/Kathmandu')
        time_ctx = self._build_telemetry_time_filter(minutes=minutes, start_time=start_time, end_time=end_time, tz=tz)
        
        cache_key = f"telemetry_{time_ctx['window_mins']}_{time_ctx.get('start_time')}_{time_ctx.get('end_time')}"
        now = time.time()
        
        # Dynamic cache TTL
        window_mins = time_ctx["window_mins"]
        if window_mins <= 5:
            cache_ttl = 5.0
        elif window_mins <= 15:
            cache_ttl = 15.0
        elif window_mins <= 60:
            cache_ttl = 30.0
        else:
            cache_ttl = 120.0
        
        if cache_key in self._cache and (now - self._cache_time.get(cache_key, 0) < cache_ttl):
            return self._cache[cache_key]

        sys_health = self.get_system_health()
        storage = self.get_storage_metrics()
        ingest = self.get_ingestion_telemetry(time_ctx)
        
        # Fast global totals query
        global_public_ips = 0
        global_subscribers = 0
        global_flows = 0
        try:
            if window_mins > 60:
                routers = ingest.get("routers", [])
                global_flows = sum(r["flows"] for r in routers) if routers else 0
                global_subscribers = sum(r["subscribers"] for r in routers) if routers else 0
                global_public_ips = sum(r["nat_ips"] for r in routers) if routers else 0
            else:
                q_global = f"""
                SELECT 
                    uniq(nat_src_ip) AS total_public_ips,
                    uniq(src_ip) AS total_subscribers,
                    count() AS total_flows
                FROM nat_logs.translations
                {time_ctx['where_sql']}
                """
                r_global = self._query(q_global, parameters=time_ctx['params']).result_rows[0]
                global_public_ips = int(r_global[0] or 0)
                global_subscribers = int(r_global[1] or 0)
                global_flows = int(r_global[2] or 0)
        except Exception as e:
            logger.error(f"Error querying global telemetry totals: {e}")
            routers = ingest.get("routers", [])
            global_flows = sum(r["flows"] for r in routers) if routers else 0

        nat_pool = self.get_public_nat_pool_details(time_ctx, limit=500, total_nat_ips_hint=global_public_ips)
        priv_pool = self.get_private_pool_details(time_ctx, limit=100, total_flows_hint=global_flows, total_subs_hint=global_subscribers)

        data = {
            "status": "ok",
            "timestamp": time.strftime("%Y-%m-%d %H:%M:%S"),
            "window": {
                "minutes": time_ctx["window_mins"],
                "seconds": time_ctx["window_sec"],
                "label": time_ctx["window_label"],
                "is_custom": time_ctx["is_custom"],
                "start_time": time_ctx["start_time"],
                "end_time": time_ctx["end_time"]
            },
            "system": sys_health,
            "storage": storage,
            "ingestion": ingest,
            "nat_pool": nat_pool,
            "private_pool": priv_pool
        }

        self._cache[cache_key] = data
        self._cache_time[cache_key] = now
        return data

    def get_public_nat_port_utilization(self, minutes=5):
        tz = getattr(config, 'SERVER_TIMEZONE', 'Asia/Kathmandu')
        time_ctx = self._build_telemetry_time_filter(minutes=minutes, tz=tz)
        return self.get_public_nat_pool_details(time_ctx, limit=500).get("top_nat_ips", [])

    def get_top_private_subscribers(self, minutes=5, limit=100):
        tz = getattr(config, 'SERVER_TIMEZONE', 'Asia/Kathmandu')
        time_ctx = self._build_telemetry_time_filter(minutes=minutes, tz=tz)
        return self.get_private_pool_details(time_ctx, limit=limit).get("top_subscribers", [])

    def get_top_destinations(self, minutes=5, limit=25):
        tz = getattr(config, 'SERVER_TIMEZONE', 'Asia/Kathmandu')
        time_ctx = self._build_telemetry_time_filter(minutes=minutes, tz=tz)
        where_sql = time_ctx["where_sql"]
        params = time_ctx["params"]
        try:
            q = f"""
            SELECT 
                dst_ip,
                count() AS flows,
                uniq(src_ip) AS subscriber_count,
                topK(1)(protocol)[1] AS dominant_proto,
                topK(1)(dst_port)[1] AS dominant_port
            FROM nat_logs.translations
            {where_sql}
            GROUP BY dst_ip
            ORDER BY flows DESC
            LIMIT {int(limit)}
            """
            res = self._query(q, parameters=params)
            destinations = []
            for r in res.result_rows:
                destinations.append({
                    "dst_ip": str(r[0]),
                    "flows": r[1],
                    "subscribers": r[2],
                    "proto": r[3],
                    "port": r[4]
                })
            return destinations
        except Exception as e:
            logger.error(f"Error querying top destinations: {e}")
            return []

    def get_active_routers(self):
        try:
            q = "SELECT DISTINCT router_ip FROM nat_logs.translations WHERE toStartOfHour(timestamp) >= toStartOfHour(now() - INTERVAL 24 HOUR) AND timestamp >= now() - INTERVAL 24 HOUR ORDER BY router_ip"
            res = self._query(q)
            routers = [str(r[0]) for r in res.result_rows if r[0]]
            return routers if routers else []
        except Exception:
            return []

    def _add_ip_or_subnet_filter(self, conditions, params, field_name, ip_value):
        if not ip_value:
            return
        val = str(ip_value).strip()
        if not val or val.upper() in ('ALL', 'NONE', '0.0.0.0/0', '0.0.0.0'):
            return
        
        try:
            if '/' in val:
                net = ipaddress.ip_network(val, strict=False)
                if net.num_addresses == 1:
                    param_key = f"{field_name}_ip"
                    conditions.append(f"{field_name} = toIPv4(%({param_key})s)")
                    params[param_key] = str(net.network_address)
                else:
                    start_key = f"{field_name}_start"
                    end_key = f"{field_name}_end"
                    conditions.append(f"{field_name} >= toIPv4(%({start_key})s) AND {field_name} <= toIPv4(%({end_key})s)")
                    params[start_key] = str(net.network_address)
                    params[end_key] = str(net.broadcast_address)
            else:
                ip_obj = ipaddress.ip_address(val)
                param_key = f"{field_name}_ip"
                conditions.append(f"{field_name} = toIPv4(%({param_key})s)")
                params[param_key] = str(ip_obj)
        except Exception:
            param_key = f"{field_name}_ip"
            conditions.append(f"{field_name} = toIPv4(%({param_key})s)")
            params[param_key] = val

    def _build_filter_clauses(self, filters):
        conditions = []
        params = {}

        start_time = filters.get('start_time')
        end_time = filters.get('end_time')
        nat_src_ip = filters.get('nat_src_ip')
        nat_src_port = filters.get('nat_src_port')
        src_ip = filters.get('src_ip')
        src_port = filters.get('src_port')
        dst_ip = filters.get('dst_ip')
        dst_port = filters.get('dst_port')
        protocol = filters.get('protocol')
        router_ip = filters.get('router_ip')

        tz = filters.get('timezone') or getattr(config, 'SERVER_TIMEZONE', 'Asia/Kathmandu')
        params['tz'] = tz

        if start_time:
            st = str(start_time).replace('T', ' ').strip()
            if len(st) == 10: st += ' 00:00:00'
            elif len(st) == 16: st += ':00'
            conditions.append(f"toStartOfHour(timestamp) >= toStartOfHour(toDateTime64(%(start_time)s, 3, '{tz}'))")
            conditions.append(f"timestamp >= toDateTime64(%(start_time)s, 3, '{tz}')")
            params['start_time'] = st

        if end_time:
            et = str(end_time).replace('T', ' ').strip()
            if len(et) == 10: et += ' 23:59:59'
            elif len(et) == 16: et += ':59'
            conditions.append(f"toStartOfHour(timestamp) <= toStartOfHour(toDateTime64(%(end_time)s, 3, '{tz}'))")
            conditions.append(f"timestamp <= toDateTime64(%(end_time)s, 3, '{tz}')")
            params['end_time'] = et

        if not start_time and not end_time and not nat_src_ip and not src_ip and not dst_ip:
            conditions.append(f"toStartOfHour(timestamp) >= toStartOfHour(now64(3, '{tz}') - INTERVAL 1 HOUR)")
            conditions.append(f"timestamp >= now64(3, '{tz}') - INTERVAL 1 HOUR")

        if nat_src_ip:
            self._add_ip_or_subnet_filter(conditions, params, 'nat_src_ip', nat_src_ip)

        if nat_src_port not in (None, '', 0, '0'):
            try:
                conditions.append("nat_src_port = %(nat_src_port)s")
                params['nat_src_port'] = int(nat_src_port)
            except ValueError:
                pass

        if src_ip:
            self._add_ip_or_subnet_filter(conditions, params, 'src_ip', src_ip)

        if src_port not in (None, '', 0, '0'):
            try:
                conditions.append("src_port = %(src_port)s")
                params['src_port'] = int(src_port)
            except ValueError:
                pass

        if dst_ip:
            self._add_ip_or_subnet_filter(conditions, params, 'dst_ip', dst_ip)

        if dst_port not in (None, '', 0, '0'):
            try:
                conditions.append("dst_port = %(dst_port)s")
                params['dst_port'] = int(dst_port)
            except ValueError:
                pass

        if protocol and str(protocol).strip().upper() not in ('ALL', '', 'NONE'):
            conditions.append("upper(protocol) = %(protocol)s")
            params['protocol'] = str(protocol).strip().upper()

        if router_ip and str(router_ip).strip().upper() not in ('ALL', '', 'NONE'):
            conditions.append("router_ip = %(router_ip)s")
            params['router_ip'] = str(router_ip).strip()

        where_clause = " WHERE " + " AND ".join(conditions) if conditions else ""
        return where_clause, params

    def query_forensics(self, filters, limit=20, with_radius=False):
        where_clause, params = self._build_filter_clauses(filters)
        tz = filters.get('timezone') or getattr(config, 'SERVER_TIMEZONE', 'Asia/Kathmandu')

        t0 = time.time()
        try:
            count_sql = f"SELECT count() FROM nat_logs.translations{where_clause}"
            count_res = self._query(count_sql, parameters=params)
            total_count = count_res.result_rows[0][0]
        except Exception as e:
            logger.error(f"Error querying forensic count: {e}")
            total_count = 0

        data_sql = f"""
        SELECT 
            toString(toTimeZone(timestamp, '{tz}')) as time_str,
            router_ip,
            protocol,
            IPv4NumToString(src_ip) as src_ip_str,
            src_port,
            IPv4NumToString(nat_src_ip) as nat_src_ip_str,
            nat_src_port,
            IPv4NumToString(dst_ip) as dst_ip_str,
            dst_port,
            toString(toTimeZone(timestamp, 'UTC')) as time_utc
        FROM nat_logs.translations
        {where_clause}
        ORDER BY timestamp DESC
        LIMIT {int(limit)}
        """
        try:
            data_res = self._query(data_sql, parameters=params)
            rows = []
            for r in data_res.result_rows:
                rows.append({
                    "timestamp": str(r[0]),
                    "timestamp_utc": str(r[9]),
                    "timezone": tz,
                    "router_ip": str(r[1]),
                    "protocol": str(r[2]),
                    "src_ip": str(r[3]),
                    "src_port": int(r[4]),
                    "nat_src_ip": str(r[5]),
                    "nat_src_port": int(r[6]),
                    "dst_ip": str(r[7]),
                    "dst_port": int(r[8]),
                    "subscriber": f"{r[3]}:{r[4]}",
                    "public_nat": f"{r[5]}:{r[6]}",
                    "destination": f"{r[7]}:{r[8]}",
                    "username": "--"
                })

            if with_radius and rows:
                try:
                    from radius_client import radius_client
                    rows = radius_client.batch_lookup_records(rows, full_details=False)
                except Exception as ex:
                    logger.warning(f"Could not enrich preview rows with RADIUS: {ex}")

        except Exception as e:
            logger.error(f"Error querying forensic data: {e}")
            rows = []

        dur_ms = round((time.time() - t0) * 1000, 1)

        return {
            "total_count": total_count,
            "query_time_ms": dur_ms,
            "returned_count": len(rows),
            "timezone": tz,
            "rows": rows
        }

    def export_forensic_records(self, filters, limit=50000, fmt="xlsx", enriched=False):
        where_clause, params = self._build_filter_clauses(filters)
        tz = filters.get('timezone') or getattr(config, 'SERVER_TIMEZONE', 'Asia/Kathmandu')

        data_sql = f"""
        SELECT 
            toString(toTimeZone(timestamp, '{tz}')) as time_str,
            toString(toTimeZone(timestamp, 'UTC')) as time_utc,
            router_ip,
            protocol,
            IPv4NumToString(src_ip) as src_ip_str,
            src_port,
            IPv4NumToString(nat_src_ip) as nat_src_ip_str,
            nat_src_port,
            IPv4NumToString(dst_ip) as dst_ip_str,
            dst_port
        FROM nat_logs.translations
        {where_clause}
        ORDER BY timestamp DESC
        LIMIT {int(limit)}
        """
        data_res = self._query(data_sql, parameters=params)
        
        records = []
        for r in data_res.result_rows:
            records.append({
                "timestamp": str(r[0]),
                "timestamp_utc": str(r[1]),
                "router_ip": str(r[2]),
                "protocol": str(r[3]),
                "src_ip": str(r[4]),
                "src_port": int(r[5]),
                "nat_src_ip": str(r[6]),
                "nat_src_port": int(r[7]),
                "dst_ip": str(r[8]),
                "dst_port": int(r[9]),
                "subscriber": f"{r[4]}:{r[5]}",
                "public_nat": f"{r[6]}:{r[7]}",
                "destination": f"{r[8]}:{r[9]}"
            })

        if enriched and records:
            try:
                from radius_client import radius_client
                records = radius_client.batch_lookup_records(records, full_details=True)
            except Exception as ex:
                logger.error(f"Could not perform RADIUS enrichment on export records: {ex}")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "NTA Lawful Interception" if enriched else "NAT Forensics"
        ws.views.sheetView[0].showGridLines = True

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        even_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        odd_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        data_font = Font(name="Calibri", size=10, color="0F172A")
        ip_font = Font(name="Consolas", size=9.5, color="0F172A")
        
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        if enriched:
            # Exact 14-Column NTA / Law Enforcement Compliance Format
            headers = [
                "PublicIP",
                "RequestDate",
                "NatIP",
                "DestinationIP",
                "StartTime",
                "StopTime",
                "InternetUsername",
                "AccountType",
                "Name",
                "PriContact",
                "SecContact",
                "PriEmail",
                "SecEmail",
                "Address"
            ]
        else:
            # Standard 13-Column Network Forensics Format
            headers = [
                f"Timestamp ({tz})",
                "Timestamp (UTC)",
                "Router Node",
                "Protocol",
                "Subscriber IP (Private)",
                "Subscriber Port",
                "Public NAT IP",
                "Public NAT Port",
                "Destination IP",
                "Destination Port",
                "Subscriber Flow (IP:Port)",
                "NAT Translation (IP:Port)",
                "Target Destination (IP:Port)"
            ]

        ws.append(headers)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws.row_dimensions[1].height = 26

        for row_idx, r in enumerate(records, start=2):
            if enriched:
                row_data = [
                    f"{r.get('nat_src_ip', '')}:{r.get('nat_src_port', '')}",
                    r.get('timestamp', ''),
                    f"{r.get('src_ip', '')}:{r.get('src_port', '')}",
                    f"{r.get('dst_ip', '')}:{r.get('dst_port', '')}",
                    r.get('session_start_time') or 'N/A',
                    r.get('session_stop_time') or 'N/A',
                    r.get('username') or 'N/A',
                    r.get('account_type') or 'Personal',
                    r.get('customer_name') or 'N/A',
                    r.get('primary_contact') or 'N/A',
                    r.get('secondary_contact') or '',
                    r.get('primary_email') or '',
                    r.get('secondary_email') or '',
                    r.get('address') or 'N/A'
                ]
            else:
                row_data = [
                    r['timestamp'],
                    r['timestamp_utc'],
                    r['router_ip'],
                    r['protocol'],
                    r['src_ip'],
                    r['src_port'],
                    r['nat_src_ip'],
                    r['nat_src_port'],
                    r['dst_ip'],
                    r['dst_port'],
                    r['subscriber'],
                    r['public_nat'],
                    r['destination']
                ]
            ws.append(row_data)
            fill_to_use = even_fill if row_idx % 2 == 0 else odd_fill

            for col_idx in range(1, len(row_data) + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                c.fill = fill_to_use
                c.border = thin_border
                c.font = ip_font if (col_idx in (1, 3, 4) if enriched else col_idx in (5, 7, 9, 11, 12, 13)) else data_font
                if col_idx in (2, 5, 6, 8, 10, 11) if enriched else col_idx in (1, 2, 4, 6, 8, 10):
                    c.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")

            ws.row_dimensions[row_idx].height = 20

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()
