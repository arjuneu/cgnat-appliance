
import io
import time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_excel(records, filter_summary=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Forensic Records"
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    even_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    odd_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    data_font = Font(name="Calibri", size=10, color="0F172A")
    ip_font = Font(name="Consolas", size=9.5, color="0F172A")
    
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    headers = [
        "Timestamp (Asia/Kathmandu)",
        "Router Node",
        "Protocol",
        "Subscriber IP (Private)",
        "Subscriber Port",
        "Public NAT IP",
        "Public NAT Port",
        "Destination IP",
        "Destination Port",
        "Subscriber Flow (IP:Port)",
        "NAT Translation (IP:Port)",
        "Target Destination (IP:Port)"
    ]

    ws.append(headers)

    # Style Header Row
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        cell.border = thin_border
    ws.row_dimensions[1].height = 26

    # Append Data Rows
    for row_idx, r in enumerate(records, start=2):
        row_data = [
            r['timestamp'],
            r['router_ip'],
            r['protocol'],
            r['src_ip'],
            r['src_port'],
            r['nat_src_ip'],
            r['nat_src_port'],
            r['dst_ip'],
            r['dst_port'],
            r['subscriber'],
            r['public_nat'],
            r['destination']
        ]
        ws.append(row_data)
        current_fill = even_fill if row_idx % 2 == 0 else odd_fill

        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = current_fill
            cell.border = thin_border
            
            # Format IP and Port cells nicely
            if col_idx in [4, 6, 8, 10, 11, 12]:
                cell.font = ip_font
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx in [5, 7, 9]:
                cell.font = data_font
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx in [2, 3]:
                cell.font = data_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.font = data_font
                cell.alignment = Alignment(horizontal="left", vertical="center")

        ws.row_dimensions[row_idx].height = 20

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()

# Test sample records
records = [
    {
        "timestamp": "2026-09-09 14:35:01.124",
        "router_ip": "203.0.113.102",
        "protocol": "TCP",
        "src_ip": "100.66.255.108",
        "src_port": 39532,
        "nat_src_ip": "203.0.113.103",
        "nat_src_port": 39532,
        "dst_ip": "103.211.149.171",
        "dst_port": 443,
        "subscriber": "100.66.255.108:39532",
        "public_nat": "203.0.113.103:39532",
        "destination": "103.211.149.171:443"
    }
]

excel_bytes = build_excel(records)
print(f"Generated beautifully styled Excel: size = {len(excel_bytes)} bytes")
