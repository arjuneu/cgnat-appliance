
import io
import time
import clickhouse_connect
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

client = clickhouse_connect.get_client(
    host='127.0.0.1',
    port=8123,
    username='default',
    password='',
    database='nat_logs'
)

def query_forensics(filters, limit=20):
    conditions = []
    params = {}

    start_time = filters.get('start_time')
    end_time = filters.get('end_time')
    nat_src_ip = filters.get('nat_src_ip')
    nat_src_port = filters.get('nat_src_port')
    src_ip = filters.get('src_ip')
    src_port = filters.get('src_port')
    dst_ip = filters.get('dst_ip')
    dst_port = filters.get('dst_port')
    protocol = filters.get('protocol')
    router_ip = filters.get('router_ip')

    if start_time:
        start_time = start_time.replace('T', ' ').strip()
        if len(start_time) == 16:
            start_time += ':00'
        conditions.append("timestamp >= toDateTime64(%(start_time)s, 3, 'Asia/Kathmandu')")
        conditions.append("toStartOfHour(timestamp) >= toStartOfHour(toDateTime64(%(start_time)s, 3, 'Asia/Kathmandu'))")
        params['start_time'] = start_time

    if end_time:
        end_time = end_time.replace('T', ' ').strip()
        if len(end_time) == 16:
            end_time += ':00'
        conditions.append("timestamp <= toDateTime64(%(end_time)s, 3, 'Asia/Kathmandu')")
        conditions.append("toStartOfHour(timestamp) <= toStartOfHour(toDateTime64(%(end_time)s, 3, 'Asia/Kathmandu'))")
        params['end_time'] = end_time

    if nat_src_ip:
        conditions.append("nat_src_ip = toIPv4(%(nat_src_ip)s)")
        params['nat_src_ip'] = nat_src_ip.strip()

    if nat_src_port:
        try:
            conditions.append("nat_src_port = %(nat_src_port)s")
            params['nat_src_port'] = int(nat_src_port)
        except ValueError:
            pass

    if src_ip:
        conditions.append("src_ip = toIPv4(%(src_ip)s)")
        params['src_ip'] = src_ip.strip()

    if src_port:
        try:
            conditions.append("src_port = %(src_port)s")
            params['src_port'] = int(src_port)
        except ValueError:
            pass

    if dst_ip:
        conditions.append("dst_ip = toIPv4(%(dst_ip)s)")
        params['dst_ip'] = dst_ip.strip()

    if dst_port:
        try:
            conditions.append("dst_port = %(dst_port)s")
            params['dst_port'] = int(dst_port)
        except ValueError:
            pass

    if protocol and protocol.upper() != 'ALL':
        conditions.append("upper(protocol) = %(protocol)s")
        params['protocol'] = protocol.strip().upper()

    if router_ip and router_ip.upper() != 'ALL':
        conditions.append("router_ip = %(router_ip)s")
        params['router_ip'] = router_ip.strip()

    where_clause = " WHERE " + " AND ".join(conditions) if conditions else ""

    t0 = time.time()
    count_sql = f"SELECT count() FROM nat_logs.translations{where_clause}"
    count_res = client.query(count_sql, parameters=params)
    total_count = count_res.result_rows[0][0]

    data_sql = f'''
    SELECT 
        toString(timestamp) as time_str,
        router_ip,
        protocol,
        IPv4NumToString(src_ip) as src_ip_str,
        src_port,
        IPv4NumToString(nat_src_ip) as nat_src_ip_str,
        nat_src_port,
        IPv4NumToString(dst_ip) as dst_ip_str,
        dst_port
    FROM nat_logs.translations
    {where_clause}
    ORDER BY timestamp DESC
    LIMIT {int(limit)}
    '''
    data_res = client.query(data_sql, parameters=params)
    dur_ms = round((time.time() - t0) * 1000, 1)

    rows = []
    for r in data_res.result_rows:
        rows.append({
            "timestamp": str(r[0]),
            "router_ip": str(r[1]),
            "protocol": str(r[2]),
            "src_ip": str(r[3]),
            "src_port": int(r[4]),
            "nat_src_ip": str(r[5]),
            "nat_src_port": int(r[6]),
            "dst_ip": str(r[7]),
            "dst_port": int(r[8]),
            "subscriber": f"{r[3]}:{r[4]}",
            "public_nat": f"{r[5]}:{r[6]}",
            "destination": f"{r[7]}:{r[8]}"
        })

    return {
        "total_count": total_count,
        "query_time_ms": dur_ms,
        "returned_count": len(rows),
        "rows": rows
    }

# Test 1: Query with IP
res1 = query_forensics({"nat_src_ip": "203.0.113.103", "limit": 5})
print("Test 1 (nat_src_ip): total =", res1['total_count'], "time_ms =", res1['query_time_ms'])
for r in res1['rows'][:2]:
    print("  Row:", r['timestamp'], r['protocol'], r['subscriber'], "->", r['public_nat'], "->", r['destination'])

# Test 2: Generate Excel buffer
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Forensic Records"
headers = [
    "Timestamp (+05:45)", "Router IP", "Protocol",
    "Subscriber IP", "Subscriber Port",
    "NAT Public IP", "NAT Public Port",
    "Destination IP", "Destination Port",
    "Subscriber Full", "NAT Public Full", "Destination Full"
]
ws.append(headers)
for r in res1['rows']:
    ws.append([
        r['timestamp'], r['router_ip'], r['protocol'],
        r['src_ip'], r['src_port'],
        r['nat_src_ip'], r['nat_src_port'],
        r['dst_ip'], r['dst_port'],
        r['subscriber'], r['public_nat'], r['destination']
    ])
buf = io.BytesIO()
wb.save(buf)
buf.seek(0)
print("Test 2 (Excel buffer): generated bytes length =", len(buf.getvalue()))
